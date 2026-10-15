"""
RFAI AI assistant — answers questions about the loaded RFAI data and builds files
(Excel, PowerPoint, chart images) on request, using Claude with tools.

The model never sees the raw table: it calls the tools below, which run on this
server against the in-memory data, and gets back only the rows/aggregates it asked for.
Engines, picked from the environment (first one set wins):
  ANTHROPIC_API_KEY  -> Claude (paid)
  GEMINI_API_KEY     -> Google Gemini free tier (no card needed)
  GROQ_API_KEY       -> Groq free tier (no card needed)
  none               -> built-in rules engine (ai_local.py): free, offline, data never leaves the server
If a free engine is over its limit or fails, the answer falls back to the built-in engine.
"""
import os, re, json, uuid, time, threading
from datetime import date, datetime
from pathlib import Path

import numpy as np
import pandas as pd

MODEL = os.environ.get('AI_MODEL', 'claude-opus-5-5')
FREE_ENGINES = {   # OpenAI-compatible chat endpoints with free tiers
    'gemini': ('https://generativelanguage.googleapis.com/v1beta/openai/chat/completions', 'gemini-flash-latest', 'Gemini'),
    'groq': ('https://api.groq.com/openai/v1/chat/completions', 'llama-3.3-70b-versatile', 'Groq'),
}
MAX_TOOL_ROUNDS = 12
_conversations = {}            # conv_id -> {"user":..., "messages":[...], "t":...}
_conv_lock = threading.Lock()
_tl = threading.local()        # PowerPoint options of the request being answered
PPT_SLIDES = {'short': 5, 'standard': 10, 'detailed': 15}
THEME_NAMES = {'airtel': 'Airtel Red', 'midnight': 'Midnight Dark', 'clean': 'Clean White', 'ocean': 'Ocean Blue'}

STATUSES = ['OnAir', 'OnAir Pending', 'Open PRI', 'Rejected RFAI', 'RFAI Pending']
GROUP_COLS = {'circle': 'circle', 'status': 'srStatus', 'category': 'category', 'project': 'bPrj',
              'toco': 'toco', 'oem': 'oem', 'domain': 'domain', 'tsp': 'tsp', 'partner': 'partner',
              'exit_month': 'exitMK', 'rfai_month': 'journeyMK', 'ageing_group': 'group'}
DATE_COLS = {'initial_rfai': 'rfaiNotice', 'rfai_rejection': 'rfaiRej', 'rfai_reoffer': 'rfaiReoffer',
             'final_rfai': 'finalRfai', 'first_pri_start': 'priFirst', 'last_pri_close': 'priLast', 'onair': 'onAir'}

SYSTEM = """You are the RFAI Tracker assistant for Airtel's 5G network deployment team.
You answer questions about the RFAI (Ready For Installation) data loaded in this tool and you
create files people ask for: Excel workbooks, PowerPoint decks and charts.

How the data works:
- One row per SR (site request). Key fields: SR number, 2G Site ID, circle, ToCo, project type,
  OEM, TSP, partner, TOCO project (category), broad project, and dates: initial RFAI, RFAI
  rejection, RFAI re-offer, final RFAI, first PRI start, last PRI close, OnAir.
- Status: OnAir, OnAir Pending, Open PRI (a Post-RFAI Issue is still open), Rejected RFAI,
  RFAI Pending. "Pending" means OnAir Pending + Open PRI + RFAI Pending.
- Days (delay): for pending SRs, days since Final RFAI with PRI days excluded; for OnAir SRs,
  days from Final RFAI to OnAir. Ageing groups: 0-30, 30-60, 60-90, >90 days.
- Exit month = OnAir month for live sites, otherwise the RFAI month.

Rules:
- Always get numbers from the tools. Never guess or invent figures. If the data can't answer
  a question, say so plainly.
- Start with get_overview when you need the available circles, months or field values.
- Prefer aggregate for counts and breakdowns; use list_srs only when the user wants specific SRs.
- When the user asks for a file (Excel, PPT, chart, report, presentation), build it with the
  file tools and mention it; the download button appears in the chat automatically.
- Keep answers short and scannable: a one-line answer first, then a compact markdown table or
  bullets if useful. Use Indian number formatting only if the user does. Today is {today}.
"""

FILTER_PROPS = {
    "circle": {"type": "string", "description": "Comma-separated circle names, e.g. 'MH,DL'"},
    "status": {"type": "string", "description": "Comma-separated statuses (OnAir, OnAir Pending, Open PRI, Rejected RFAI, RFAI Pending) or 'pending' for all not-yet-OnAir"},
    "category": {"type": "string", "description": "Comma-separated TOCO project / category values"},
    "project": {"type": "string", "description": "Comma-separated broad project values"},
    "toco": {"type": "string"}, "oem": {"type": "string"}, "domain": {"type": "string", "description": "Project type"},
    "month": {"type": "string", "description": "Comma-separated exit months like \"Jan'26,Feb'26\""},
    "aging": {"type": "string", "enum": ["0-30", "30-60", "60-90", ">90", ">30"]},
    "search": {"type": "string", "description": "Part of an SR number or site ID"},
    "min_days": {"type": "integer", "description": "Only SRs with more than this many days (delay)"},
    "max_days": {"type": "integer", "description": "Only SRs with fewer than this many days (delay)"},
    "date_field": {"type": "string", "enum": list(DATE_COLS), "description": "Restrict to SRs whose date in this field falls between date_from and date_to"},
    "date_from": {"type": "string", "description": "YYYY-MM-DD"},
    "date_to": {"type": "string", "description": "YYYY-MM-DD"},
}

def _tools():
    flt = {"type": "object", "properties": FILTER_PROPS}
    return [
        {"name": "get_overview",
         "description": "Totals, status mix, the list of circles, exit months, categories, projects, OEMs and ToCos in the loaded data, and the data's date range. Call first when unsure what values exist.",
         "input_schema": {"type": "object", "properties": {}}},
        {"name": "aggregate",
         "description": "Count SRs (and average/max days) grouped by up to two fields, with optional filters. Use for any 'how many', 'by circle', 'per month', 'breakdown', 'top N' question.",
         "input_schema": {"type": "object", "properties": {
             "group_by": {"type": "array", "items": {"type": "string", "enum": list(GROUP_COLS)}, "maxItems": 2,
                          "description": "Fields to group by; empty for a single total"},
             "date_group": {"type": "object", "description": "Group instead by a date field per day/week/month",
                            "properties": {"field": {"type": "string", "enum": list(DATE_COLS)},
                                           "period": {"type": "string", "enum": ["day", "week", "month"]}}},
             "filters": flt,
             "top": {"type": "integer", "description": "Keep only the largest N groups"}}}},
        {"name": "list_srs",
         "description": "List individual SRs with their key fields, filtered and sorted. Returns at most 100 rows plus the total match count.",
         "input_schema": {"type": "object", "properties": {
             "filters": flt,
             "sort_by": {"type": "string", "enum": ["days", "sr", "circle", "initial_rfai", "final_rfai", "onair"]},
             "descending": {"type": "boolean"}, "limit": {"type": "integer", "maximum": 100}}}},
        {"name": "get_sr",
         "description": "Everything known about one SR (all fields and its journey dates).",
         "input_schema": {"type": "object", "properties": {"sr": {"type": "string"}}, "required": ["sr"]}},
        {"name": "make_data_excel",
         "description": "Build the standard formatted RFAI Excel workbook (summary sheet with charts, Journey Table, pending / PRI / rejected / OnAir lists) for SRs matching the filters.",
         "input_schema": {"type": "object", "properties": {"filters": flt, "title": {"type": "string"}}, "required": ["title"]}},
        {"name": "make_custom_excel",
         "description": "Build a formatted Excel workbook from tables you supply (e.g. aggregate results). Each sheet: name, columns, rows; optional chart.",
         "input_schema": {"type": "object", "properties": {
             "title": {"type": "string"},
             "sheets": {"type": "array", "items": {"type": "object", "properties": {
                 "name": {"type": "string"}, "columns": {"type": "array", "items": {"type": "string"}},
                 "rows": {"type": "array", "items": {"type": "array"}},
                 "chart": {"type": "object", "properties": {"type": {"type": "string", "enum": ["column", "bar", "line", "pie"]},
                                                            "title": {"type": "string"}}}},
                 "required": ["name", "columns", "rows"]}}},
             "required": ["title", "sheets"]}},
        {"name": "make_powerpoint",
         "description": "Build a polished Airtel-styled PowerPoint deck (cover and closing slides are added automatically). Slide kinds: 'title' (title, subtitle, up to 3 headline kpis), 'summary' (executive summary: stats:[{value,label}] up to 3 big numbers, points:[{head,text}] up to 3 findings), 'kpis' (title, kpis:[{label,value,sub}]), 'stat' (one big number: value, label, text:[2-3 lines of context]), 'chart' (title, chart_type, categories, series:[{name,values}], insights), 'table' (title, columns, rows; add a 'Total' row; long tables continue on extra slides), 'bullets' (title, bullets), 'section' (title, subtitle), 'closing' (dark next-steps slide: bullets). Write each slide title as the takeaway (e.g. 'MH holds 26% of pending SRs', no full stop), put context in subtitle, give charts 2-4 short insights and every slide speaker notes. A good review deck: title, summary, kpis, 2-4 charts, 1-2 tables, closing with next steps.",
         "input_schema": {"type": "object", "properties": {
             "title": {"type": "string"},
             "slides": {"type": "array", "items": {"type": "object", "properties": {
                 "kind": {"type": "string", "enum": ["title", "summary", "kpis", "stat", "chart", "table", "bullets", "section", "closing"]},
                 "stats": {"type": "array", "items": {"type": "object", "properties": {"value": {"type": "string"}, "label": {"type": "string"}}}},
                 "points": {"type": "array", "items": {"type": "object", "properties": {"head": {"type": "string"}, "text": {"type": "string"}}}},
                 "value": {"type": "string"}, "label": {"type": "string"}, "text": {"type": "array", "items": {"type": "string"}},
                 "title": {"type": "string"}, "subtitle": {"type": "string"},
                 "bullets": {"type": "array", "items": {"type": "string"}},
                 "kpis": {"type": "array", "items": {"type": "object", "properties": {"label": {"type": "string"}, "value": {"type": "string"}, "sub": {"type": "string"}}}},
                 "insights": {"type": "array", "items": {"type": "string"}},
                 "columns": {"type": "array", "items": {"type": "string"}}, "rows": {"type": "array", "items": {"type": "array"}},
                 "chart_type": {"type": "string", "enum": ["column", "bar", "line", "pie", "stacked_column"]},
                 "categories": {"type": "array", "items": {"type": "string"}},
                 "series": {"type": "array", "items": {"type": "object", "properties": {"name": {"type": "string"}, "values": {"type": "array", "items": {"type": "number"}}}}},
                 "notes": {"type": "string"}},
                 "required": ["kind", "title"]}}},
             "required": ["title", "slides"]}},
        {"name": "make_chart_image",
         "description": "Draw a chart as a PNG image the user can download or paste anywhere.",
         "input_schema": {"type": "object", "properties": {
             "title": {"type": "string"}, "chart_type": {"type": "string", "enum": ["column", "bar", "line", "pie", "stacked_column"]},
             "categories": {"type": "array", "items": {"type": "string"}},
             "series": {"type": "array", "items": {"type": "object", "properties": {"name": {"type": "string"}, "values": {"type": "array", "items": {"type": "number"}}}}}},
             "required": ["title", "chart_type", "categories", "series"]}},
    ]


def _loose(schema):
    """Schema for OpenAI-compatible engines: drop limits some providers reject, give every array its items."""
    if isinstance(schema, dict):
        out = {k: _loose(v) for k, v in schema.items() if k not in ('maxItems', 'maximum')}
        if out.get('type') == 'array' and not out.get('items'): out['items'] = {"type": "string"}
        if out.get('type') == 'object' and 'properties' in out and not out['properties']: out.pop('properties')
        return out
    if isinstance(schema, list): return [_loose(x) for x in schema]
    return schema


def _num(v):
    if isinstance(v, str) and re.fullmatch(r'-?\d[\d,]*(\.\d+)?', v.strip()):
        try: return float(v.replace(',', '')) if '.' in v else int(v.replace(',', ''))
        except ValueError: return v
    return v


def _xlsx_preview(path, max_sheets=8, max_rows=30, max_cols=14):
    """First rows of each sheet, for the chat's spreadsheet preview."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheets = []
        for ws in wb.worksheets[:max_sheets]:
            rows = []
            for r in ws.iter_rows(max_row=max_rows, max_col=max_cols, values_only=True):
                rows.append(['' if v is None else (v.strftime('%d-%b-%Y') if hasattr(v, 'strftime') else
                             (round(v, 2) if isinstance(v, float) else v)) for v in r])
            while rows and not any(str(v).strip() for v in rows[-1]): rows.pop()
            sheets.append({"name": ws.title, "rows": rows, "total": ws.max_row})
        wb.close()
        return {"type": "sheets", "sheets": sheets}
    except Exception:
        return None


ST_COL = {'onair': '#10B981', 'onair pending': '#F59E0B', 'pending': '#F59E0B', 'open pri': '#3B82F6',
          'rejected rfai': '#EF4444', 'rejected': '#EF4444', 'rfai pending': '#94A3B8'}
def _col(name, i, pal):
    return ST_COL.get(str(name).strip().lower(), pal[i % len(pal)])


class Assistant:
    """Bound to the host app: get_df() returns the current data, build_xlsx is the export builder."""

    def __init__(self, get_df, xl_filter, build_xlsx, mk_sort_val, files_dir):
        self.get_df, self.xl_filter, self.build_xlsx, self.mk_sort_val = get_df, xl_filter, build_xlsx, mk_sort_val
        self.files_dir = Path(files_dir); self.files_dir.mkdir(parents=True, exist_ok=True)
        self.files = {}           # file_id -> {"path", "name", "user", "t"}
        self._client = None

    # ── public ──────────────────────────────────────────────────────────────
    def enabled(self):
        return True

    def engine(self):
        if os.environ.get('ANTHROPIC_API_KEY') or os.environ.get('ANTHROPIC_AUTH_TOKEN'): return 'claude'
        if os.environ.get('GEMINI_API_KEY') or os.environ.get('GOOGLE_API_KEY'): return 'gemini'
        if os.environ.get('GROQ_API_KEY'): return 'groq'
        return 'builtin'

    def engine_label(self):
        return {'claude': 'Claude AI', 'gemini': 'Gemini AI · free', 'groq': 'Groq AI · free'}.get(self.engine(), 'Built-in · free & offline')

    def chat(self, user, conv_id, text, ppt=None):
        """Returns (conv_id, answer_markdown, files, suggestions, views).
        views: data behind the answer (last breakdown or SR list), drawn as interactive charts in the chat.
        ppt: optional PowerPoint choices {length: short|standard|detailed, theme, images: none|art|ai}."""
        with _conv_lock:
            self._gc()
            conv = _conversations.get(conv_id) if conv_id else None
            if not conv or conv['user'] != user:
                conv_id = uuid.uuid4().hex; conv = {'user': user, 'messages': [], 'omsgs': [], 'last': None, 't': time.time()}
                _conversations[conv_id] = conv
        conv['t'] = time.time()
        if ppt:
            conv['ppt'] = {'length': ppt.get('length') if ppt.get('length') in PPT_SLIDES else 'standard',
                           'theme': ppt.get('theme') if ppt.get('theme') in THEME_NAMES else 'airtel',
                           'images': ppt.get('images') if ppt.get('images') in ('none', 'art', 'ai') else 'none'}
        _tl.ppt = conv.get('ppt') or {'length': 'standard', 'theme': 'airtel', 'images': 'none'}
        _tl.calls = []
        eng = self.engine()
        if ppt and eng != 'builtin':
            text += f"\n\n(PowerPoint settings chosen by the user: about {PPT_SLIDES[_tl.ppt['length']]} slides. Theme and pictures are applied automatically.)"
        if eng == 'builtin':
            ans, files, sug = self._local(conv, text, user)
            return conv_id, ans, files, sug, self._views(files)
        try:
            if eng == 'claude':
                ans, files = self._chat_claude(conv, text, user)
            else:
                ans, files = self._chat_openai(conv, text, user, eng)
            return conv_id, ans, files, [], self._views(files)
        except Exception as e:
            if eng == 'claude': raise
            _tl.calls = []
            ans, files, sug = self._local(conv, text, user)
            why = 'is over its free limit right now' if '429' in str(e) else 'did not answer'
            return conv_id, f"_{FREE_ENGINES[eng][2]} {why}, so the built-in assistant answered._\n\n" + ans, files, sug, self._views(files)

    # explorer filter names for the chat's "open these SRs" clicks
    _EV = {'initial_rfai': 'rfai', 'final_rfai': 'final', 'first_pri_start': 'pri_open', 'last_pri_close': 'pri_close',
           'onair': 'onair', 'rfai_rejection': 'rej'}

    def _xp_filters(self, f):
        out = {k: str(v) for k, v in (f or {}).items() if v not in (None, '') and k in FILTER_PROPS and k not in ('date_field', 'date_from', 'date_to', 'search')}
        f = f or {}
        if f.get('date_field') in self._EV:
            out['ev'] = self._EV[f['date_field']]
            if f.get('date_from'): out['dfrom'] = str(f['date_from'])
            if f.get('date_to'): out['dto'] = str(f['date_to'])
        if f.get('search'): out['q'] = str(f['search'])
        out.pop('min_days', None); out.pop('max_days', None)
        if f.get('min_days') and not out.get('aging'):
            n = int(f['min_days']); out['aging'] = '>90' if n >= 90 else ('>30' if n >= 30 else '')
            if not out['aging']: out.pop('aging')
        return out

    def _views(self, files):
        """The data behind the answer, for an interactive chart (skipped when a file with its own preview was made)."""
        if files: return []
        for name, a, out in reversed(getattr(_tl, 'calls', [])):
            if not isinstance(out, dict) or out.get('error'): continue
            flt = self._xp_filters(a.get('filters'))
            if name == 'aggregate' and out.get('rows'):
                rows = out['rows'][:60]
                dims = [k for k in rows[0] if k not in ('srs', 'avg_days', 'max_days')]
                dg = a.get('date_group') or {}
                return [{'type': 'agg', 'dims': dims, 'rows': rows, 'total': out.get('total_srs'), 'filters': flt,
                         'date': {'ev': self._EV.get(dg.get('field')), 'period': dg.get('period', 'month')} if dg.get('field') else None}]
            if name == 'aggregate' and out.get('statuses') and len(out['statuses']) > 1:
                rows = [{'status': k, 'srs': v} for k, v in out['statuses'].items()]
                return [{'type': 'agg', 'dims': ['status'], 'rows': rows, 'total': out.get('total'), 'filters': flt, 'date': None}]
            if name == 'list_srs' and out.get('rows'):
                return [{'type': 'srs', 'rows': out['rows'][:25], 'total': out.get('total_matches'), 'filters': flt}]
        return []

    def _local(self, conv, text, user):
        from ai_local import LocalBrain
        ans, files, state, sug = LocalBrain(self).answer(text, user, conv.get('last'), getattr(_tl, 'ppt', None))
        if state is not None: conv['last'] = state
        return ans, files, sug

    def _chat_openai(self, conv, text, user, eng):
        """Tool loop against an OpenAI-compatible endpoint (Gemini / Groq free tiers)."""
        import requests
        url, model, _ = FREE_ENGINES[eng]
        key = (os.environ.get('GEMINI_API_KEY') or os.environ.get('GOOGLE_API_KEY')) if eng == 'gemini' else os.environ.get('GROQ_API_KEY')
        model = os.environ.get('AI_FREE_MODEL') or model
        tools = [{"type": "function", "function": {"name": t['name'], "description": t['description'],
                  **({"parameters": _loose(t['input_schema'])} if t['input_schema'].get('properties') else {})}} for t in _tools()]
        msgs = conv['omsgs']; start = len(msgs)
        msgs.append({"role": "user", "content": text}); made = []
        sysmsg = {"role": "system", "content": SYSTEM.replace('{today}', date.today().strftime('%d %b %Y'))}
        try:
            for _ in range(MAX_TOOL_ROUNDS):
                r = requests.post(url, headers={"Authorization": f"Bearer {key}", "Content-Type": "application/json"},
                                  json={"model": model, "messages": [sysmsg] + msgs[-40:], "tools": tools, "temperature": 0.2}, timeout=120)
                if r.status_code >= 400:
                    raise RuntimeError(f"{r.status_code}: {r.text[:300]}")
                m = r.json()['choices'][0]['message']
                calls = m.get('tool_calls') or []
                msgs.append({"role": "assistant", "content": m.get('content') or '', **({"tool_calls": calls} if calls else {})})
                if not calls:
                    return (m.get('content') or 'Done.').strip(), made
                for c in calls:
                    try:
                        args = json.loads(c['function'].get('arguments') or '{}') or {}
                        out = self._run_tool(c['function']['name'], args, user)
                        if isinstance(out, dict) and out.get('_file'):
                            made.append(out['_file']); out = {k: v for k, v in out.items() if k != '_file'}
                        content = json.dumps(out, default=str)[:30000]
                    except Exception as e:
                        content = f"Error: {e}"
                    msgs.append({"role": "tool", "tool_call_id": c.get('id', ''), "name": c['function']['name'], "content": content})
            return "That took too many steps. Try a more specific question.", made
        except Exception:
            del msgs[start:]   # keep the history valid
            raise

    def _chat_claude(self, conv, text, user):
        import anthropic
        if self._client is None:
            self._client = anthropic.Anthropic(max_retries=3)
        msgs = conv['messages']; msgs.append({"role": "user", "content": text})
        made = []; tools = _tools()
        system = [{"type": "text", "text": SYSTEM.replace('{today}', date.today().strftime('%d %b %Y')),
                   "cache_control": {"type": "ephemeral"}}]
        for _ in range(MAX_TOOL_ROUNDS):
            resp = self._client.beta.messages.create(
                model=MODEL, max_tokens=16000, system=system, tools=tools, messages=msgs,
                output_config={"effort": "medium"},
                betas=["server-side-fallback-2026-07-01"], fallbacks="default",
            )
            if resp.stop_reason == "refusal":
                msgs.pop()   # keep history valid for the next question
                return "Sorry, I can't help with that request.", made
            msgs.append({"role": "assistant", "content": resp.content})
            if resp.stop_reason != "tool_use":
                answer = "\n".join(b.text for b in resp.content if b.type == "text").strip()
                return answer or "Done.", made
            results = []
            for b in resp.content:
                if b.type != "tool_use": continue
                try:
                    out = self._run_tool(b.name, b.input or {}, user)
                    if isinstance(out, dict) and out.get('_file'):
                        made.append(out['_file'])
                        out = {k: v for k, v in out.items() if k != '_file'}
                    results.append({"type": "tool_result", "tool_use_id": b.id, "content": json.dumps(out, default=str)[:60000]})
                except Exception as e:
                    results.append({"type": "tool_result", "tool_use_id": b.id, "content": f"Error: {e}", "is_error": True})
            msgs.append({"role": "user", "content": results})
        return "That took too many steps. Try a more specific question.", made

    def file(self, file_id, user):
        f = self.files.get(file_id)
        if not f or f['user'] != user or not os.path.exists(f['path']): return None
        return f

    # ── tools ───────────────────────────────────────────────────────────────
    def _df(self, filters=None):
        df = self.get_df()
        if df is None or not len(df): raise RuntimeError("No data is loaded yet.")
        p = {k: str(v) for k, v in (filters or {}).items() if v not in (None, '') and k in FILTER_PROPS}
        if p.get('date_field'):
            p['ev'] = {'initial_rfai': 'rfai', 'final_rfai': 'final', 'first_pri_start': 'pri_open',
                       'last_pri_close': 'pri_close', 'onair': 'onair', 'rfai_rejection': 'rej'}.get(p.pop('date_field'), '')
            if not p['ev']: p.pop('ev')
            if 'date_from' in p: p['dfrom'] = p.pop('date_from')
            if 'date_to' in p: p['dto'] = p.pop('date_to')
        else:
            p.pop('date_from', None); p.pop('date_to', None)
        lo, hi = p.pop('min_days', None), p.pop('max_days', None)
        df = self.xl_filter(df, p)
        if lo or hi:                       # any day threshold, e.g. "older than 120 days"
            d = pd.to_numeric(df['delay'], errors='coerce')
            if lo: df = df[d > float(lo)]
            if hi: df = df[d < float(hi)]
        return df

    @staticmethod
    def _days(df):
        return pd.to_numeric(df['delay'], errors='coerce').where(lambda x: x >= 0)

    def _run_tool(self, name, a, user):
        out = self._run_tool_(name, a, user)
        if name in ('aggregate', 'list_srs') and hasattr(_tl, 'calls'): _tl.calls.append((name, a, out))
        return out

    def _run_tool_(self, name, a, user):
        if name == 'get_overview':
            df = self.get_df()
            if df is None or not len(df): return {"error": "No data is loaded yet."}
            st = df['srStatus'].astype(str).value_counts().to_dict()
            vals = lambda c: sorted(x for x in df[c].astype(str).unique() if x and x != 'nan')[:80] if c in df.columns else []
            d = pd.to_datetime(df['rfaiNotice'].astype(str), format='%d-%b-%Y', errors='coerce')
            return {"total_srs": len(df), "statuses": st, "circles": vals('circle'),
                    "exit_months": sorted(vals('exitMK'), key=self.mk_sort_val), "categories": vals('category'),
                    "projects": vals('bPrj'), "oems": vals('oem'), "tocos": vals('toco'), "project_types": vals('domain'),
                    "initial_rfai_from": str(d.min().date()) if d.notna().any() else None,
                    "initial_rfai_to": str(d.max().date()) if d.notna().any() else None}
        if name == 'aggregate':
            df = self._df(a.get('filters'))
            days = self._days(df)
            keys = []
            dg = a.get('date_group') or {}
            if dg.get('field'):
                d = pd.to_datetime(df[DATE_COLS[dg['field']]].astype(str), format='%d-%b-%Y', errors='coerce')
                per = dg.get('period', 'month')
                lab = d.dt.strftime('%Y-%m-%d') if per == 'day' else (d - pd.to_timedelta(d.dt.weekday, unit='D')).dt.strftime('week of %Y-%m-%d') if per == 'week' else d.dt.strftime('%Y-%m')
                keys.append(lab.rename(dg['field'] + '_' + per))
            for g in (a.get('group_by') or [])[:2]:
                keys.append(df[GROUP_COLS[g]].astype(str).rename(g))
            if not keys:
                return {"total": len(df), "avg_days": None if days.isna().all() else round(float(days.mean()), 1),
                        "statuses": df['srStatus'].astype(str).value_counts().to_dict()}
            g = pd.DataFrame({k.name: k for k in keys}).assign(_d=days)
            out = g.groupby([k.name for k in keys], dropna=True).agg(srs=('_d', 'size'), avg_days=('_d', 'mean'), max_days=('_d', 'max')).reset_index()
            out = out.sort_values('srs', ascending=False)
            if dg.get('field'): out = out.sort_values(keys[0].name)
            if a.get('top'): out = out.nlargest(int(a['top']), 'srs')
            out['avg_days'] = out['avg_days'].round(1)
            rows = out.head(300).replace({np.nan: None}).to_dict('records')
            return {"total_srs": len(df), "groups": len(out), "rows": rows, "truncated": len(out) > 300}
        if name == 'list_srs':
            df = self._df(a.get('filters'))
            sk = {'days': None, 'sr': 'srNo', 'circle': 'circle', 'initial_rfai': 'rfaiNotice', 'final_rfai': 'finalRfai', 'onair': 'onAir'}.get(a.get('sort_by') or 'days')
            desc = a.get('descending', True)
            if sk is None: key = self._days(df)
            elif sk in ('rfaiNotice', 'finalRfai', 'onAir'): key = pd.to_datetime(df[sk].astype(str), format='%d-%b-%Y', errors='coerce')
            else: key = df[sk].astype(str)
            df = df.loc[key.sort_values(ascending=not desc, na_position='last').index]
            n = min(int(a.get('limit') or 25), 100)
            cols = ['srNo', 'siteId', 'circle', 'srStatus', 'delay', 'bPrj', 'category', 'oem', 'rfaiNotice', 'finalRfai', 'priFirst', 'priLast', 'onAir']
            sub = df[[c for c in cols if c in df.columns]].head(n).astype(str).replace({'nan': '', 'None': '', '-1': ''})
            return {"total_matches": len(df), "shown": len(sub), "rows": sub.to_dict('records')}
        if name == 'get_sr':
            df = self.get_df()
            hit = df[df['srNo'].astype(str).str.upper() == str(a['sr']).strip().upper()]
            if not len(hit): hit = df[df['srNo'].astype(str).str.upper().str.contains(str(a['sr']).strip().upper(), regex=False)].head(5)
            if not len(hit): return {"error": "SR not found"}
            return {"matches": hit.astype(str).replace({'nan': ''}).head(5).to_dict('records')}
        if name == 'make_data_excel':
            df = self._df(a.get('filters')).reset_index(drop=True)
            if not len(df): return {"error": "No SRs match these filters."}
            ftxt = ' · '.join(f"{k}: {v}" for k, v in (a.get('filters') or {}).items() if v)
            path, fid, fname = self._new_file(a.get('title') or 'RFAI export', '.xlsx', user)
            self.build_xlsx(df, path, ftxt, a.get('title') or 'RFAI export', include_journey=True)
            return {"ok": True, "srs": len(df), "file": fname, "_file": self._meta(fid, _xlsx_preview(path))}
        if name == 'make_custom_excel':
            path, fid, fname = self._new_file(a['title'], '.xlsx', user)
            self._custom_xlsx(path, a['title'], a['sheets'])
            return {"ok": True, "file": fname, "_file": self._meta(fid, _xlsx_preview(path))}
        if name == 'make_powerpoint':
            path, fid, fname = self._new_file(a['title'], '.pptx', user)
            prev = self._pptx(path, a['title'], a['slides'], user)
            return {"ok": True, "file": fname, "slides": len(a['slides']), "_file": self._meta(fid, prev)}
        if name == 'make_chart_image':
            path, fid, fname = self._new_file(a['title'], '.png', user)
            self._png(path, a)
            return {"ok": True, "file": fname, "_file": self._meta(fid, {"type": "image"})}
        return {"error": f"Unknown tool {name}"}

    # ── files ───────────────────────────────────────────────────────────────
    def _new_file(self, title, ext, user):
        fid = uuid.uuid4().hex
        fname = 'RFAI_' + re.sub(r'[^A-Za-z0-9]+', '_', title)[:50].strip('_') + '_' + datetime.now().strftime('%d%b%Y_%H%M') + ext
        path = str(self.files_dir / (fid + ext))
        self.files[fid] = {"path": path, "name": fname, "user": user, "t": time.time()}
        return path, fid, fname

    def _meta(self, fid, preview=None):
        f = self.files[fid]; m = {"id": fid, "name": f['name'], "kind": f['name'].rsplit('.', 1)[-1]}
        if preview: m['preview'] = preview
        return m

    def _gc(self):
        now = time.time()
        for k in [k for k, v in _conversations.items() if now - v['t'] > 6 * 3600]: _conversations.pop(k, None)
        for k in [k for k, v in self.files.items() if now - v['t'] > 6 * 3600]:
            try:
                if not self.files[k].get('keep'): os.remove(self.files[k]['path'])   # shared deck pictures stay cached
            except Exception: pass
            self.files.pop(k, None)

    def _custom_xlsx(self, path, title, sheets):
        import xlsxwriter
        wb = xlsxwriter.Workbook(path)
        F = lambda **k: wb.add_format({'font_name': 'Calibri', 'font_size': 10, 'valign': 'vcenter', **k})
        ban = F(bold=True, font_size=16, font_color='#FFFFFF', bg_color='#E40000', indent=1)
        hdr = F(bold=True, font_color='#FFFFFF', bg_color='#2B2F3A', border=1, border_color='#2B2F3A', align='center', text_wrap=True)
        txt = F(border=1, border_color='#E6D9DB'); num = F(border=1, border_color='#E6D9DB', num_format='#,##0.##', align='center')
        used = set()
        for sh in sheets[:20]:
            nm = re.sub(r'[\[\]:*?/\\]', ' ', sh.get('name') or 'Sheet')[:31] or 'Sheet'
            while nm in used: nm = nm[:28] + str(len(used))
            used.add(nm)
            ws = wb.add_worksheet(nm); ws.hide_gridlines(2)
            cols = sh['columns']; rows = [[_num(v) for v in r] for r in sh['rows']]
            ws.set_row(0, 30); ws.merge_range(0, 0, 0, max(1, len(cols) - 1), '  ' + title + (' — ' + sh['name'] if sh.get('name') else ''), ban)
            ws.write_row(2, 0, cols, hdr); ws.set_row(2, 28)
            for i, r in enumerate(rows):
                for j, v in enumerate(r[:len(cols)]):
                    if isinstance(v, (int, float)) and not isinstance(v, bool): ws.write_number(3 + i, j, v, num)
                    else: ws.write(3 + i, j, '' if v is None else str(v), txt)
            for j, c in enumerate(cols):
                w = max([len(str(c))] + [len(str(r[j])) for r in rows[:200] if j < len(r)]) + 3
                ws.set_column(j, j, min(max(w, 9), 45))
            ws.freeze_panes(3, 1)
            if rows: ws.autofilter(2, 0, 2 + len(rows), len(cols) - 1)
            chp = sh.get('chart')
            if chp and rows and len(cols) > 1:
                ch = wb.add_chart({'type': chp.get('type', 'column')})
                pal = ['#E40000', '#10B981', '#F59E0B', '#3B82F6', '#94A3B8', '#A855F7']
                for j in range(1, len(cols)):
                    if not all(isinstance(r[j], (int, float)) for r in rows if j < len(r) and r[j] is not None): continue
                    ser = {'name': [nm, 2, j], 'categories': [nm, 3, 0, 2 + len(rows), 0], 'values': [nm, 3, j, 2 + len(rows), j]}
                    if chp.get('type') != 'pie': ser['fill'] = {'color': pal[(j - 1) % len(pal)]}
                    ch.add_series(ser)
                ch.set_title({'name': chp.get('title') or sh['name']}); ch.set_size({'width': 720, 'height': 360})
                ws.insert_chart(2, len(cols) + 1, ch)
        wb.close()

    def _pptx(self, path, title, slides, user=''):
        """Builds the deck; returns preview data (the final slide list) for the chat's slide viewer."""
        from ppt_builder import build_deck
        o = getattr(_tl, 'ppt', None) or {}
        theme = o.get('theme', 'airtel')
        built, imgs = build_deck(path, title, slides, theme, o.get('images', 'none'))
        pics = {}
        for slot, p in (imgs or {}).items():       # pictures are served through the same per-user file endpoint
            if p and os.path.exists(p):
                fid = uuid.uuid4().hex
                self.files[fid] = {"path": p, "name": f'{slot}{os.path.splitext(p)[1]}', "user": user, "t": time.time(), "keep": True}
                pics[slot] = fid
        keys = ('kind', 'title', 'subtitle', 'kpis', 'stats', 'points', 'value', 'label', 'text', 'chart_type', 'categories',
                'series', 'insights', 'columns', 'rows', 'bullets')
        out = []
        for s in built:
            x = {k: s[k] for k in keys if k in s}
            if 'rows' in x: x['rows'] = [[None if v is None else (v if isinstance(v, (int, float)) else str(v)) for v in r] for r in x['rows'][:13]]
            if 'categories' in x: x['categories'] = [str(c) for c in x['categories'][:30]]
            out.append(x)
        return {"type": "slides", "theme": theme, "slides": out, "pics": pics}

    def _png(self, path, a):
        """Chart picture in the app's AI style: gradient rounded bars, highlighted leader, glow line with shaded area, donut."""
        import matplotlib; matplotlib.use('Agg')
        import matplotlib.pyplot as plt
        from matplotlib.patches import FancyBboxPatch
        from matplotlib.colors import LinearSegmentedColormap, to_rgb
        pal = ['#E40000', '#10B981', '#F59E0B', '#3B82F6', '#94A3B8', '#A855F7']
        cats = [str(c) for c in a['categories']]; ser = a['series']; ct = a['chart_type']
        vals = [[float(_num(t) or 0) if not isinstance(t, (int, float)) else float(t) for t in s_['values']][:len(cats)] for s_ in ser]
        INK, MUT, GRID = '#16181F', '#8A8F9A', '#EEF0F4'
        plt.rcParams.update({'font.family': 'DejaVu Sans', 'font.size': 11})
        fig = plt.figure(figsize=(12, 6.6), dpi=160, facecolor='white')
        ax = fig.add_axes([0.07, 0.13, 0.9, 0.62])

        def shade(hex_, k):
            r, g, b = to_rgb(hex_)
            return tuple(min(1, c + (1 - c) * k) if k > 0 else c * (1 + k) for c in (r, g, b))

        def grad_patch(x, y, w, h, color, horizontal=False, glow=False):
            if w <= 0 or h <= 0: return
            r = min(w, h) * 0.18 if not horizontal else min(w, h) * 0.18
            pt = FancyBboxPatch((x, y), w, h, boxstyle=f'round,pad=0,rounding_size={r}', linewidth=0, facecolor='none', mutation_aspect=1)
            ax.add_patch(pt)
            cm = LinearSegmentedColormap.from_list('g', [shade(color, -0.12), shade(color, 0.32)])
            img = ax.imshow(np.linspace(0, 1, 64).reshape(1, -1) if horizontal else np.linspace(0, 1, 64).reshape(-1, 1),
                            cmap=cm, extent=(x, x + w, y, y + h), origin='lower', aspect='auto', zorder=3)
            img.set_clip_path(pt)
            if glow and not horizontal:
                for k, al in ((1.06, .10), (1.12, .05)):
                    ax.add_patch(FancyBboxPatch((x - w * (k - 1) / 2, y), w * k, h * (1 + (k - 1) / 3), boxstyle=f'round,pad=0,rounding_size={r}',
                                                linewidth=0, facecolor=color, alpha=al, zorder=2))

        total = sum(sum(v) for v in vals) or 1
        if ct == 'pie':
            ax.set_position([0.05, 0.08, 0.6, 0.7])
            v = vals[0]; cols = [_col(c, i, pal) for i, c in enumerate(cats)]
            w, _ = ax.pie(v, colors=cols, startangle=90, counterclock=False, wedgeprops={'width': 0.36, 'edgecolor': 'white', 'linewidth': 3})
            for wd, x in zip(w, v):
                if x / total > .04:
                    ang = np.deg2rad((wd.theta1 + wd.theta2) / 2)
                    ax.text(0.82 * np.cos(ang), 0.82 * np.sin(ang), f'{x / total:.0%}', ha='center', va='center', color='white', fontsize=13, fontweight='bold')
            ax.text(0, 0.06, f'{int(total):,}', ha='center', va='center', fontsize=30, fontweight='bold', color=INK)
            ax.text(0, -0.16, 'TOTAL SRS', ha='center', va='center', fontsize=10, fontweight='bold', color=MUT)
            ax.set_aspect('equal')
            lx = fig.add_axes([0.66, 0.15, 0.3, 0.55]); lx.axis('off')
            for i, (c, x) in enumerate(sorted(zip(cats, v), key=lambda p: -p[1])[:10]):
                y = 0.95 - i * 0.1
                lx.add_patch(FancyBboxPatch((0, y - 0.03), 0.05, 0.06, boxstyle='round,pad=0,rounding_size=0.01', color=_col(c, cats.index(c), pal), transform=lx.transAxes))
                lx.text(0.09, y, c, va='center', fontsize=12, color=INK, transform=lx.transAxes)
                lx.text(0.98, y, f'{int(x):,}  ·  {x / total:.0%}', va='center', ha='right', fontsize=12, color=MUT, fontweight='bold', transform=lx.transAxes)
        else:
            n = len(cats); multi = len(vals) > 1
            stacked = ct == 'stacked_column'; horiz = ct == 'bar'; line = ct == 'line'
            top = int(np.argmax(vals[0])) if vals and vals[0] else 0
            mx = max([sum(col) for col in zip(*vals)] if stacked else [max(v) for v in vals] + [1]) * 1.15
            for sp in ax.spines.values(): sp.set_visible(False)
            ax.tick_params(length=0, colors=MUT)
            if line:
                xs = np.arange(n)
                for si, v in enumerate(vals):
                    c = _col(ser[si]['name'], si, pal) if multi else '#E40000'
                    for lw, al in ((14, .05), (9, .08), (5, .15)): ax.plot(xs, v, color=c, linewidth=lw, alpha=al, solid_capstyle='round', zorder=2)
                    ax.plot(xs, v, color=c, linewidth=3.2, zorder=4, solid_capstyle='round')
                    if not multi:
                        poly = ax.fill_between(xs, v, 0, color='none')
                        grad = ax.imshow(np.linspace(0, 1, 64).reshape(-1, 1), cmap=LinearSegmentedColormap.from_list('a', [(1, 1, 1, 0), (*to_rgb(c), .35)]),
                                         extent=(xs[0], xs[-1], 0, max(v) or 1), origin='lower', aspect='auto', zorder=1)
                        grad.set_clip_path(poly.get_paths()[0], transform=ax.transData)
                    ax.scatter(xs, v, s=[140 if (i == top and not multi) else 60 for i in range(n)], color='white', edgecolor=c, linewidth=2.6, zorder=5)
                    if n <= 16:
                        for i, y in enumerate(v):
                            ax.annotate(f'{y:,.0f}', (i, y), textcoords='offset points', xytext=(0, 11), ha='center', fontsize=10,
                                        fontweight='bold', color='#C4000F' if (i == top and not multi) else '#4B505C', zorder=6)
                ax.set_xticks(xs, cats); ax.set_xlim(-0.5, n - 0.5); ax.set_ylim(0, mx)
                ax.grid(axis='y', color=GRID, linewidth=1)
            else:
                k = 1 if stacked else len(vals); slot = 1.0; bw = 0.72 / k
                acc = np.zeros(n)
                for si, v in enumerate(vals):
                    for i, y in enumerate(v):
                        hi = not multi and i == top and n > 1
                        c = _col(ser[si]['name'], si, pal) if multi else ('#E40000' if hi or n == 1 else '#B9C0CC')
                        pos = i - 0.36 + (0 if stacked else si * bw)
                        if horiz:
                            grad_patch(acc[i] if stacked else 0, n - 1 - i - 0.36 + (0 if stacked else si * bw), y, bw, c, True, hi)
                            if not stacked: ax.text(y + mx * .01, n - 1 - i - 0.36 + si * bw + bw / 2, f'{y:,.0f}', va='center', fontsize=10, fontweight='bold', color='#C4000F' if hi else INK, zorder=6)
                        else:
                            grad_patch(pos, acc[i] if stacked else 0, bw, y, c, False, hi)
                            if not stacked and n <= 24: ax.text(pos + bw / 2, y + mx * .015, f'{y:,.0f}', ha='center', fontsize=10, fontweight='bold', color='#C4000F' if hi else INK, zorder=6)
                            if stacked and y > mx * .05: ax.text(pos + bw / 2, acc[i] + y / 2, f'{y:,.0f}', ha='center', va='center', fontsize=9, fontweight='bold', color='white', zorder=6)
                    if stacked: acc += np.array(v)
                if horiz:
                    ax.set_ylim(-0.6, n - 0.4); ax.set_xlim(0, mx); ax.set_yticks(range(n), list(reversed(cats)))
                    ax.grid(axis='x', color=GRID, linewidth=1)
                else:
                    ax.set_xlim(-0.6, n - 0.4); ax.set_ylim(0, mx); ax.set_xticks(range(n), cats)
                    ax.grid(axis='y', color=GRID, linewidth=1)
            ax.set_axisbelow(True)
            ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda x, _: f'{x:,.0f}')) if not horiz else \
                ax.xaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda x, _: f'{x:,.0f}'))
            plt.setp(ax.get_xticklabels(), rotation=30 if n > 10 and not horiz else 0, ha='right' if n > 10 and not horiz else 'center', color='#4B505C')
            plt.setp(ax.get_yticklabels(), color='#4B505C' if horiz else MUT)
            if multi:
                for si, s_ in enumerate(ser):
                    fig.text(0.07 + si * 0.14, 0.785, '●  ' + str(s_['name']), color=_col(s_['name'], si, pal), fontsize=11, fontweight='bold')
        # header: title, an insight line and a small AI badge
        fig.text(0.07, 0.93, a['title'], fontsize=20, fontweight='bold', color=INK)
        v0 = [sum(c) for c in zip(*vals)] if ct == 'stacked_column' and vals else (vals[0] if vals else [])
        if v0 and ct != 'line':
            i = int(np.argmax(v0)); ins = f'✦  {cats[i]} leads with {v0[i]:,.0f} ({v0[i] / (sum(v0) or 1):.0%} of total)'
        elif v0:
            i = int(np.argmax(v0)); ins = f'✦  Peak in {cats[i]} at {v0[i]:,.0f} · {sum(v0):,.0f} in total'
        else: ins = ''
        fig.text(0.07, 0.875, ins, fontsize=12, color='#6B4FD8', fontweight='bold')
        fig.text(0.97, 0.935, 'AI INSIGHT', fontsize=9, fontweight='bold', color='white', ha='right',
                 bbox=dict(boxstyle='round,pad=0.45,rounding_size=0.8', fc='#E40000', ec='none'))
        fig.text(0.07, 0.03, 'RFAI Tracker  ·  ' + datetime.now().strftime('%d %b %Y, %H:%M'), fontsize=9, color=MUT)
        lp = Path(__file__).with_name('assets') / 'airtel_red.png'
        if lp.exists():
            lim = plt.imread(str(lp)); la = fig.add_axes([0.875, 0.015, 0.095, 0.05]); la.imshow(lim); la.axis('off')
        fig.savefig(path, facecolor='white'); plt.close(fig)
